"""Turn uploaded files into something a model can read: base64 images for
vision models, plain text for everything else."""

import base64
import binascii
import io
import logging
import mimetypes
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Tuple

from .config import Config

log = logging.getLogger(__name__)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}
# Ollama decodes PNG and JPEG reliably; other formats are converted when Pillow is present.
NATIVE_IMAGE_EXT = {".png", ".jpg", ".jpeg"}
TEXT_EXT = {
    ".txt", ".md", ".rst", ".log", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".html", ".htm", ".css", ".js", ".ts", ".tsx", ".jsx",
    ".py", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".java", ".kt", ".rs", ".go", ".php", ".rb",
    ".sh", ".bat", ".ps1", ".sql", ".tex", ".bib", ".ui", ".ipynb",
}


class AttachmentError(ValueError):
    pass


@dataclass
class Attachment:
    name: str
    kind: str              # "image" or "text"
    data: str              # base64 for images, extracted text otherwise
    truncated: bool = False


def decode_data_url(url: str) -> Tuple[Optional[str], bytes]:
    """Split `data:<mime>;base64,<payload>` (or bare base64) into mime and bytes."""
    mime = None
    payload = url
    if url.startswith("data:"):
        header, _, payload = url.partition(",")
        mime = header[5:].split(";")[0] or None
    try:
        raw = base64.b64decode(payload, validate=False)
    except (binascii.Error, ValueError) as e:
        raise AttachmentError(f"bad base64 data: {e}")
    if len(raw) > Config.MAX_ATTACHMENT_BYTES:
        raise AttachmentError(f"attachment is {len(raw)} bytes; the limit is {Config.MAX_ATTACHMENT_BYTES}")
    return mime, raw


def _ext(name: str, mime: Optional[str]) -> str:
    ext = Path(name or "").suffix.lower()
    if not ext and mime:
        ext = mimetypes.guess_extension(mime) or ""
    return ext


def image_to_base64(raw: bytes, ext: str) -> str:
    if ext not in NATIVE_IMAGE_EXT:
        try:
            from PIL import Image
            with Image.open(io.BytesIO(raw)) as img:
                buf = io.BytesIO()
                img.convert("RGB").save(buf, format="PNG")
                raw = buf.getvalue()
        except ImportError:
            log.warning("Pillow not installed; sending %s image unconverted", ext)
        except Exception as e:
            raise AttachmentError(f"could not read image: {e}")
    return base64.b64encode(raw).decode("ascii")


def _decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _looks_textual(raw: bytes) -> bool:
    sample = raw[:4096]
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def _pdf_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise AttachmentError("PDF support needs `pip install pypdf`")
    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[page {i}]\n{text}")
    if not pages:
        return "[This PDF has no extractable text; it may be a scan.]"
    return "\n\n".join(pages)


def _docx_text(raw: bytes) -> str:
    try:
        import docx
    except ImportError:
        raise AttachmentError("Word support needs `pip install python-docx`")
    doc = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _html_text(raw: bytes) -> str:
    from .search import html_to_text
    return html_to_text(_decode_text(raw))


def extract_text(name: str, raw: bytes, mime: Optional[str] = None) -> str:
    ext = _ext(name, mime)
    if ext == ".pdf" or mime == "application/pdf":
        text = _pdf_text(raw)
    elif ext == ".docx":
        text = _docx_text(raw)
    elif ext in (".html", ".htm"):
        text = _html_text(raw)
    elif ext in TEXT_EXT or (mime or "").startswith("text/") or _looks_textual(raw):
        text = _decode_text(raw)
    else:
        raise AttachmentError(f"{name}: file type {ext or mime or 'unknown'} is not supported")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def load_attachment(name: str, raw: bytes, mime: Optional[str] = None) -> Attachment:
    name = name or "attachment"
    ext = _ext(name, mime)
    if ext in IMAGE_EXT or (mime or "").startswith("image/"):
        return Attachment(name, "image", image_to_base64(raw, ext))
    text = extract_text(name, raw, mime)
    truncated = len(text) > Config.FILE_TEXT_CHARS
    if truncated:
        text = text[: Config.FILE_TEXT_CHARS]
    return Attachment(name, "text", text, truncated)


def attachment_block(att: Attachment) -> str:
    note = " (truncated)" if att.truncated else ""
    return f'<file name="{att.name}"{note}>\n{att.data}\n</file>'
